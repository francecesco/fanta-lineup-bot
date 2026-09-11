"""Mappa nome giocatore ↔ ID fantacalcio.it, con ruolo e squadra.

Due fonti, stessa interfaccia (`resolve`, `role_of`, `owned`):

- `RosterMap`: legge dai file Excel — listone `Quotazioni_*.xlsx` (colonna `Id`)
  per nome→ID/ruolo e `rosa.xlsx` per i giocatori posseduti.
- `RosterSito`: ricava tutto dai dati del sito (`get_lineup` → `lineUpInfo`,
  campi `pid`/`plyr`/`role`), senza alcun file Excel. È la fonte usata dal bot
  autonomo: bastano le credenziali del sito.

Uso tipico:
    rm = RosterMap()               # oppure RosterSito(res_get_lineup)
    rm.resolve("Calhanoglu")   -> (2194, "C")
"""

import os
import unicodedata

import config

# openpyxl è importato in modo "lazy" dentro i metodi di caricamento: così build_payload
# può essere usato con un RosterMap alternativo anche senza openpyxl installato.

# Ruolo dall'API del sito (lineUpInfo[].role): 1=Portiere, 2=Difensore, 3=Centrocampista, 4=Attaccante.
RUOLO_API = {1: "P", 2: "D", 3: "C", 4: "A"}


def _norm(name):
    """Normalizza un nome per il confronto tollerante (accenti, maiuscole, spazi)."""
    s = unicodedata.normalize("NFKD", str(name))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().split())


def _role_api(role):
    """Ruolo (P/D/C/A) dal campo `role` di lineUpInfo, che è una lista es. [1]."""
    if isinstance(role, list):
        role = role[0] if role else None
    try:
        return RUOLO_API.get(int(role), "")
    except (TypeError, ValueError):
        return ""


def _risolvi(by_name, by_norm, owned, name, require_owned, fonte):
    """Logica di risoluzione nome→(id, ruolo) condivisa dalle due fonti."""
    key = str(name).strip()
    if require_owned and key not in owned and _norm(key) not in {_norm(k) for k in owned}:
        raise ValueError(f"'{name}' non è nella tua rosa")
    if key in by_name:
        pid, role, _ = by_name[key]
        return pid, role
    cands = by_norm.get(_norm(key), [])
    if len(cands) == 1:
        _n, pid, role = cands[0]
        return pid, role
    if not cands:
        raise ValueError(f"'{name}' non trovato in {fonte}")
    names = ", ".join(n for n, _, _ in cands)
    raise ValueError(f"'{name}' è ambiguo in {fonte}: {names}. Usa il nome esatto.")


class _BaseRoster:
    """Interfaccia comune: struttura dati (`_by_name`, `_by_norm`, `owned`) + risoluzione."""

    _fonte = "la rosa"

    def resolve(self, name, require_owned=True):
        """Restituisce (id, ruolo) per un giocatore. Solleva ValueError se non risolvibile.

        Con require_owned=True, il giocatore deve essere posseduto (in rosa).
        """
        return _risolvi(self._by_name, self._by_norm, self.owned, name, require_owned, self._fonte)

    def role_of(self, name):
        return self.resolve(name)[1]

    def aggiorna(self, res):
        """Aggiorna la rosa dai dati del sito. Le fonti statiche (file) non fanno nulla."""
        return None


class RosterMap(_BaseRoster):
    _fonte = "il listone quotazioni"

    def __init__(self, quotazioni=None, rosa=None, base_dir="."):
        self.quotazioni_path = os.path.join(base_dir, quotazioni or config.FILE_QUOTAZIONI)
        self.rosa_path = os.path.join(base_dir, rosa or config.FILE_ROSA)
        self._by_name = {}       # nome esatto -> (id, ruolo, squadra)
        self._by_norm = {}       # nome normalizzato -> lista di (nome, id, ruolo)
        self.owned = {}          # nome esatto in rosa -> ruolo
        self._load_quotazioni()
        self._load_rosa()

    # ---- caricamento ----
    def _load_quotazioni(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.quotazioni_path, read_only=True, data_only=True)
        ws = wb["Tutti"] if "Tutti" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = next(i for i, r in enumerate(rows[:5])
                     if r and any(str(c).strip().lower() == "id" for c in r if c is not None))
        hdr = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_i]]
        ci, cn, cr = hdr.index("id"), hdr.index("nome"), hdr.index("r")
        cs = hdr.index("squadra") if "squadra" in hdr else None
        for r in rows[hdr_i + 1:]:
            if r[ci] is None or r[cn] is None:
                continue
            name = str(r[cn]).strip()
            pid = int(r[ci])
            role = str(r[cr]).strip().upper()
            team = str(r[cs]).strip() if cs is not None and r[cs] is not None else ""
            self._by_name[name] = (pid, role, team)
            self._by_norm.setdefault(_norm(name), []).append((name, pid, role))

    def _load_rosa(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.rosa_path, read_only=True, data_only=True)
        ws = wb["Rosa"] if "Rosa" in wb.sheetnames else wb.active
        rows = [r for r in ws.iter_rows(values_only=True)
                if any(c is not None and str(c).strip() for c in r)]
        hdr = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        ir = hdr.index("ruolo") if "ruolo" in hdr else 0
        ig = hdr.index("giocatore") if "giocatore" in hdr else 1
        rolemap = {"P": "P", "D": "D", "C": "C", "A": "A",
                   "POR": "P", "DIF": "D", "CEN": "C", "ATT": "A"}
        for r in rows[1:]:
            if ig >= len(r) or r[ig] is None:
                continue
            name = str(r[ig]).strip()
            role = rolemap.get(str(r[ir]).strip().upper(), "") if ir < len(r) and r[ir] else ""
            if name:
                self.owned[name] = role


class RosterSito(_BaseRoster):
    """Rosa ricavata dai dati del sito (get_lineup → lineUpInfo), senza file Excel.

    `lineUpInfo` è la tua rosa completa da 25: ogni elemento porta `pid` (ID API),
    `plyr` (nome) e `role` (ruolo). Passa la risposta di get_lineup al costruttore,
    oppure chiama `aggiorna(res)` per rinfrescarla a ogni lettura dal sito.
    """

    _fonte = "i dati del sito"

    def __init__(self, res=None):
        self._by_name = {}
        self._by_norm = {}
        self.owned = {}
        if res is not None:
            self.aggiorna(res)

    def aggiorna(self, res):
        by_name, by_norm, owned = {}, {}, {}
        for p in (res.get("lineUpInfo") or []):
            name = str(p.get("plyr") or "").strip()
            pid = p.get("pid")
            if not name or pid is None:
                continue
            role = _role_api(p.get("role"))
            team = str(p.get("tname") or "")
            by_name[name] = (int(pid), role, team)
            by_norm.setdefault(_norm(name), []).append((name, int(pid), role))
            owned[name] = role
        # Riassegnazione atomica: una resolve concorrente vede sempre una rosa completa.
        self._by_name, self._by_norm, self.owned = by_name, by_norm, owned
